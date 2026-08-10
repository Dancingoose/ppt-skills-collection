import importlib.util
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook


SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "inventory_material.py"
SPEC = importlib.util.spec_from_file_location("inventory_material", SCRIPT)
INVENTORY = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(INVENTORY)


class InventoryMaterialTests(unittest.TestCase):
    def temporary_file(self, name, content, binary=False):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        path = Path(directory.name) / name
        if binary:
            path.write_bytes(content)
        else:
            path.write_text(content, encoding="utf-8")
        return path

    def test_markdown_text_is_preserved(self):
        source = self.temporary_file("brief.md", "# Spring fair\n\n60+ clubs")
        text, warnings = INVENTORY.extract(source)
        self.assertEqual(text, "# Spring fair\n\n60+ clubs")
        self.assertEqual(warnings, [])

    def test_docx_paragraphs_are_extracted_in_order(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        source = Path(directory.name) / "plan.docx"
        document = """<?xml version='1.0' encoding='UTF-8'?>
        <w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>
          <w:body>
            <w:p><w:r><w:t>Spring fair</w:t></w:r></w:p>
            <w:p><w:r><w:t>60+ clubs</w:t></w:r></w:p>
          </w:body>
        </w:document>"""
        with zipfile.ZipFile(source, "w") as archive:
            archive.writestr("word/document.xml", document)
        text, warnings = INVENTORY.extract(source)
        self.assertEqual(text, "Spring fair\n\n60+ clubs")
        self.assertEqual(warnings, [])

    def test_unsupported_material_fails_explicitly(self):
        source = self.temporary_file("brief.xyz", "opaque")
        text, warnings = INVENTORY.extract(source)
        self.assertEqual(text, "")
        self.assertEqual(warnings, ["Unsupported source type: .xyz"])

    def test_standalone_image_is_registered_for_visual_review(self):
        source = self.temporary_file("poster.png", b"not-a-real-image", binary=True)
        refs = INVENTORY.standalone_image_refs(source)
        self.assertEqual(refs[0]["name"], "poster.png")
        self.assertEqual(refs[0]["source"], str(source.resolve()))
        self.assertIn("inspect visually", refs[0]["description"])

    def test_csv_keeps_headers_rows_and_numeric_columns(self):
        source = self.temporary_file("budget.csv", "项目,金额(元),说明\n宣传物料,2500,\"海报, 横幅\"\n舞台设备,4000,音响\n")
        tables, warnings = INVENTORY.csv_table_refs(source)
        self.assertEqual(warnings, [])
        self.assertEqual(tables[0]["headers"], ["项目", "金额(元)", "说明"])
        self.assertEqual(tables[0]["rowCount"], 2)
        self.assertEqual(tables[0]["numericColumns"], [{"name": "金额(元)", "values": [2500.0, 4000.0]}])

    def test_xlsx_keeps_each_sheet_and_flags_uncalculated_formulas(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        source = Path(directory.name) / "budget.xlsx"
        workbook = Workbook()
        costs = workbook.active
        costs.title = "Costs"
        costs.append(["Workstream", "Budget"])
        costs.append(["Venue", 2500])
        costs.append(["Production", 4000])
        summary = workbook.create_sheet("Summary")
        summary.append(["Metric", "Value"])
        summary.append(["Total", "=SUM(Costs!B2:B3)"])
        workbook.save(source)
        workbook.close()

        text, warnings = INVENTORY.extract(source)
        tables, table_warnings = INVENTORY.xlsx_table_refs(source)

        self.assertIn("## Worksheet: Costs", text)
        self.assertIn("Venue\t2500", text)
        self.assertEqual(warnings, [])
        self.assertEqual(tables[0]["sheet"], "Costs")
        self.assertEqual(tables[0]["rowCount"], 2)
        self.assertEqual(tables[0]["numericColumns"], [{"name": "Budget", "values": [2500.0, 4000.0]}])
        self.assertEqual(tables[1]["formulaCells"], 1)
        self.assertIn("without cached results", table_warnings[0])

    def test_xlsx_uses_a_cached_formula_result_as_numeric_evidence(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        source = Path(directory.name) / "cached.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Metric", "Value"])
        sheet.append(["Total", "=SUM(2500,4000)"])
        workbook.save(source)
        workbook.close()

        with zipfile.ZipFile(source) as archive:
            members = {info.filename: archive.read(info.filename) for info in archive.infolist()}
        worksheet = members["xl/worksheets/sheet1.xml"].decode("utf-8")
        worksheet = worksheet.replace(
            "<f>SUM(2500,4000)</f><v></v>",
            "<f>SUM(2500,4000)</f><v>6500</v>",
        )
        members["xl/worksheets/sheet1.xml"] = worksheet.encode("utf-8")
        with zipfile.ZipFile(source, "w") as archive:
            for name, data in members.items():
                archive.writestr(name, data)

        tables, warnings = INVENTORY.xlsx_table_refs(source)

        self.assertEqual(warnings, [])
        self.assertEqual(tables[0]["numericColumns"], [{"name": "Value", "values": [6500.0]}])

    def test_web_fetch_saves_raw_html_and_strips_script_content(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)

        class Response:
            headers = SimpleNamespace(get_content_type=lambda: "text/html", get_content_charset=lambda: "utf-8")
            def read(self, _): return b"<html><body><h1>Event brief</h1><p>60 clubs</p><script>ignore me</script></body></html>"
            def __enter__(self): return self
            def __exit__(self, *_): return False

        with patch.object(INVENTORY, "urlopen", return_value=Response()):
            text, images, warnings, saved = INVENTORY.fetch_web_material("https://example.com/brief", Path(directory.name))
        self.assertIn("Event brief", text)
        self.assertIn("60 clubs", text)
        self.assertNotIn("ignore me", text)
        self.assertEqual(images, [])
        self.assertEqual(warnings, [])
        self.assertTrue(saved.is_file())

    def test_web_fetch_rejects_non_http_url(self):
        text, images, warnings, saved = INVENTORY.fetch_web_material("file:///secret.html", Path(tempfile.gettempdir()))
        self.assertEqual((text, images, saved), ("", [], None))
        self.assertIn("http or https", warnings[0])

    def test_pdf_text_keeps_page_boundaries(self):
        source = self.temporary_file("brief.pdf", b"not parsed by the mock", binary=True)

        class Reader:
            pages = [
                SimpleNamespace(extract_text=lambda: "Cover claim"),
                SimpleNamespace(extract_text=lambda: "60+ clubs"),
            ]

        with patch.dict(sys.modules, {"pypdf": SimpleNamespace(PdfReader=lambda _: Reader())}):
            text, warnings = INVENTORY.extract(source)
        self.assertEqual(text, "## Page 1\nCover claim\n\n## Page 2\n60+ clubs")
        self.assertEqual(warnings, [])

    def test_image_only_pdf_requires_visual_or_ocr_review(self):
        source = self.temporary_file("poster.pdf", b"not parsed by the mock", binary=True)

        class Reader:
            pages = [SimpleNamespace(extract_text=lambda: "")]

        with patch.dict(sys.modules, {"pypdf": SimpleNamespace(PdfReader=lambda _: Reader())}):
            text, warnings = INVENTORY.extract(source)
        self.assertEqual(text, "## Page 1\n[No extractable text]")
        self.assertIn("inspect visually or use OCR", warnings[0])

    def test_pdf_images_are_written_as_visual_evidence(self):
        source = self.temporary_file("poster.pdf", b"not parsed by the mock", binary=True)
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)

        class Reader:
            pages = [SimpleNamespace(images=[SimpleNamespace(name="poster.png", data=b"image-bytes")])]

        with patch.dict(sys.modules, {"pypdf": SimpleNamespace(PdfReader=lambda _: Reader())}):
            refs = INVENTORY.pdf_image_refs(source, Path(directory.name))
        self.assertEqual(refs[0]["page"], 1)
        self.assertTrue(Path(refs[0]["source"]).is_file())
        self.assertEqual(Path(refs[0]["source"]).read_bytes(), b"image-bytes")


if __name__ == "__main__":
    unittest.main()
