#!/usr/bin/env python3
"""Extract supported source material into auditable PPT workflow inventory files."""
import argparse
import json
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET


WORD_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def plain_text(path: Path):
    return path.read_text(encoding="utf-8", errors="replace"), []


def docx_text(path: Path):
    warnings = []
    try:
        with zipfile.ZipFile(path) as archive:
            root = ET.fromstring(archive.read("word/document.xml"))
    except (OSError, KeyError, zipfile.BadZipFile, ET.ParseError) as exc:
        return "", [f"DOCX extraction failed: {exc}"]
    paragraphs = []
    for paragraph in root.iter(f"{WORD_NS}p"):
        text = "".join(node.text or "" for node in paragraph.iter(f"{WORD_NS}t"))
        if text.strip():
            paragraphs.append(text.strip())
    return "\n\n".join(paragraphs), warnings


def pptx_text(path: Path):
    warnings = []
    try:
        from pptx import Presentation
    except ImportError:
        return "", ["python-pptx is unavailable; run with the workflow virtual environment."]
    try:
        presentation = Presentation(path)
    except Exception as exc:
        return "", [f"PPTX extraction failed: {exc}"]
    chunks, images = [], []
    for slide_index, slide in enumerate(presentation.slides, 1):
        lines = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                value = shape.text.strip()
                if value:
                    lines.append(value)
            if getattr(shape, "shape_type", None) == 13:  # MSO_SHAPE_TYPE.PICTURE
                images.append({"slide": slide_index, "name": shape.name, "source": "embedded picture"})
        chunks.append(f"## Slide {slide_index}\n" + ("\n".join(lines) or "[No extractable text]"))
    return "\n\n".join(chunks), warnings + ([f"Found {len(images)} embedded pictures."] if images else [])


def pptx_image_refs(path: Path):
    try:
        from pptx import Presentation
        presentation = Presentation(path)
    except Exception:
        return []
    refs = []
    for slide_index, slide in enumerate(presentation.slides, 1):
        for shape in slide.shapes:
            if getattr(shape, "shape_type", None) == 13:
                refs.append({"slide": slide_index, "name": shape.name, "description": "Embedded picture"})
    return refs


def extract(path: Path):
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".html", ".htm", ".csv"}:
        return plain_text(path)
    if suffix == ".docx":
        return docx_text(path)
    if suffix == ".pptx":
        return pptx_text(path)
    return "", [f"Unsupported source type: {suffix or 'no extension'}"]


def markdown_escape(value: str):
    return value.replace("\r\n", "\n").strip()


def main():
    parser = argparse.ArgumentParser(description="Create auditable material inventory for a PPT task")
    parser.add_argument("source", type=Path)
    parser.add_argument("--task", required=True, type=Path)
    parser.add_argument("--method", default="file extraction")
    args = parser.parse_args()
    if not args.source.is_file():
        raise SystemExit(f"Source file does not exist: {args.source}")
    args.task.mkdir(parents=True, exist_ok=True)
    text, warnings = extract(args.source)
    image_refs = pptx_image_refs(args.source) if args.source.suffix.lower() == ".pptx" else []
    inventory = {
        "source": str(args.source.resolve()), "method": args.method,
        "format": args.source.suffix.lower(), "text": text, "images": image_refs,
        "warnings": warnings,
    }
    (args.task / "material-inventory.json").write_text(json.dumps(inventory, ensure_ascii=False, indent=2), encoding="utf-8")
    warning_lines = "\n".join(f"- {warning}" for warning in warnings) or "- None"
    image_lines = "\n".join(f"- Slide {item['slide']}: {item['name']} ({item['description']})" for item in image_refs) or "- None detected"
    content = f"""# Content Inventory

## Source Files
- File: {args.source.resolve()}
- Extraction method: {args.method}
- Structured inventory: material-inventory.json

## Extracted Source Content
{markdown_escape(text) or '[No text could be extracted. Review the original material manually.]'}

## Core Message
[To be confirmed from the material with the creator.]

## Image Resources
{image_lines}

## Data Points
[Extract claims with source URLs or citations before using a data layout.]

## Page and Chapter Plan
[Create after confirming the audience, decision, and core message.]

## Extraction Warnings
{warning_lines}
"""
    (args.task / "content-inventory.md").write_text(content, encoding="utf-8")
    print(f"Wrote {args.task / 'content-inventory.md'}")
    print(f"Wrote {args.task / 'material-inventory.json'}")
    if warnings:
        print("Warnings:")
        for warning in warnings:
            print(f"- {warning}")


if __name__ == "__main__":
    main()
