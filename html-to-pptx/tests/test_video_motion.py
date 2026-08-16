import os
import sys
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from pptx import Presentation


SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))
sys.path.insert(0, str(SCRIPTS.parent))

import convert as converter  # noqa: E402
import video_motion  # noqa: E402


class VideoMotionTests(unittest.TestCase):
    def test_ffmpeg_command_uses_powerpoint_compatible_h264(self):
        command = video_motion._ffmpeg_command("ffmpeg", Path("frames"), 20, Path("clip.mp4"))
        self.assertIn("libx264", command)
        self.assertIn("yuv420p", command)
        self.assertIn("+faststart", command)
        self.assertEqual(command[command.index("-framerate") + 1], "20")

    @unittest.skipUnless(os.name == "nt", "PowerPoint media embedding is Windows-only")
    def test_records_and_embeds_marked_slide_video(self):
        if not video_motion.ffmpeg_executable():
            self.skipTest("FFmpeg unavailable; set PPT_FFMPEG_EXECUTABLE to run media integration")
        try:
            import win32com.client  # noqa: F401
        except ImportError as exc:
            self.skipTest(f"pywin32 unavailable: {exc}")

        html = """<!doctype html><style>
        .slide { width: 640px; height: 360px; position: relative; overflow: hidden; background: #10243d; }
        .dot { position: absolute; width: 64px; height: 64px; border-radius: 50%; background: #f4b942;
               animation: move 400ms linear forwards; }
        @keyframes move { from { transform: translate(0, 150px); } to { transform: translate(500px, 150px); } }
        </style><section class='slide' data-pptx-slide data-pptx-video
        data-pptx-video-duration='400' data-pptx-video-fps='10'><div class='dot'></div></section>"""
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "animated.html"
            media_dir = root / "media"
            output = root / "video.pptx"
            source.write_text(html, encoding="utf-8")
            videos = video_motion.record_marked_slide_videos(source, media_dir)
            self.assertEqual(len(videos), 1)
            self.assertGreater(Path(videos[0]["path"]).stat().st_size, 1024)

            presentation = Presentation()
            presentation.slides.add_slide(presentation.slide_layouts[6])
            presentation.save(output)
            self.assertEqual(video_motion.embed_slide_videos(output, videos), [1])
            with ZipFile(output) as archive:
                self.assertTrue(any(name.endswith(".mp4") and name.startswith("ppt/media/")
                                    for name in archive.namelist()))

    @unittest.skipUnless(os.name == "nt", "PowerPoint media embedding is Windows-only")
    def test_convert_can_embed_marked_slide_video(self):
        if not video_motion.ffmpeg_executable():
            self.skipTest("FFmpeg unavailable; set PPT_FFMPEG_EXECUTABLE to run media integration")
        try:
            import win32com.client  # noqa: F401
        except ImportError as exc:
            self.skipTest(f"pywin32 unavailable: {exc}")
        html = """<!doctype html><style>
        .slide { width: 640px; height: 360px; position: relative; background: #132238; }
        .bar { width: 50px; height: 100%; background: #56c7b8; animation: grow 300ms linear forwards; }
        @keyframes grow { from { transform: scaleX(.1); transform-origin: left; } to { transform: scaleX(1); transform-origin: left; } }
        </style><section class='slide' data-pptx-slide data-pptx-video
        data-pptx-video-duration='300' data-pptx-video-fps='10'><div class='bar'></div></section>"""
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "deck.html"
            output = root / "deck.pptx"
            source.write_text(html, encoding="utf-8")
            converter.convert(
                source, output, keep_screenshots=False, embed_fonts=False,
                do_verify=False, do_preflight=False, do_visual_audit=False,
                embed_video_motion=True,
            )
            with ZipFile(output) as archive:
                self.assertTrue(any(name.endswith(".mp4") and name.startswith("ppt/media/")
                                    for name in archive.namelist()))


if __name__ == "__main__":
    unittest.main()
