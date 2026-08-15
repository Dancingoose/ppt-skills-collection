import importlib.util
import os
import sys
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE


SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))


def load_module(name):
    spec = importlib.util.spec_from_file_location(name, SCRIPTS / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


MOTION = load_module("motion")
MEASURE = load_module("measure")
ASSEMBLE = load_module("assemble")
MOTION_AUDIT = load_module("motion_audit")


class MotionTests(unittest.TestCase):
    def test_emits_composable_motion_plan_and_triggers(self):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, 0, 0, 1000000, 1000000)
        plan = [
            {"effect": "fade", "duration": 300, "trigger": "click"},
            {"effect": "motion", "path": "M 0 0 L 0.18 -0.08 E", "duration": 650, "trigger": "with"},
            {"effect": "scale", "from": [80, 80], "to": [115, 115], "duration": 650, "trigger": "with"},
            {"effect": "rotate", "by": 360, "duration": 650, "trigger": "with"},
            {"effect": "fade", "duration": 250, "trigger": "after"},
        ]
        result = MOTION.apply_native_animations(
            slide, [{"id": "hero", "plan": plan}], {"hero": shape}
        )
        self.assertEqual(result["emitted"], ["hero"])
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "complex-motion.pptx"
            presentation.save(output)
            with ZipFile(output) as archive:
                root = etree.fromstring(archive.read("ppt/slides/slide1.xml"))
        ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
        self.assertEqual(root.xpath("count(.//p:animMotion[@path='M 0 0 L 0.18 -0.08 E'])", namespaces=ns), 1.0)
        self.assertEqual(root.xpath("count(.//p:animScale/p:from[@x='800' and @y='800'])", namespaces=ns), 1.0)
        self.assertEqual(root.xpath("count(.//p:animRot[@by='21600000'])", namespaces=ns), 1.0)
        self.assertEqual(root.xpath("count(.//p:cTn[@nodeType='withEffect'])", namespaces=ns), 3.0)
        self.assertEqual(root.xpath("count(.//p:cTn[@nodeType='afterEffect'])", namespaces=ns), 1.0)

    def test_emits_native_timing_for_fade_and_spin(self):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        fade_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, 1000000, 1000000)
        spin_shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, 1000000, 0, 1000000, 1000000)

        result = MOTION.apply_native_animations(slide, [
            {"id": "fade", "kind": "fade", "duration": 600, "delay": 0},
            {"id": "spin", "kind": "spin", "duration": 800, "delay": 120},
        ], {"fade": fade_shape, "spin": spin_shape})

        self.assertEqual(result["emitted"], ["fade", "spin"])
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "motion.pptx"
            presentation.save(output)
            with ZipFile(output) as archive:
                root = etree.fromstring(archive.read("ppt/slides/slide1.xml"))
        ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
        self.assertEqual(len(root.xpath("./p:timing", namespaces=ns)), 1)
        self.assertEqual(root.xpath("count(.//p:animEffect[@filter='fade'])", namespaces=ns), 1.0)
        self.assertEqual(root.xpath("count(.//p:spTgt/p:bg)", namespaces=ns), 0.0)
        self.assertEqual(root.xpath("count(.//p:animRot[@by='21600000'])", namespaces=ns), 1.0)
        self.assertEqual({int(v) for v in root.xpath(".//p:bldP/@spid", namespaces=ns)},
                         {fade_shape.shape_id, spin_shape.shape_id})

    def test_uses_native_text_build_without_background_only_flag(self):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, 1000000, 1000000)
        text = slide.shapes.add_textbox(100000, 100000, 800000, 400000)
        text.text_frame.text = "Animated text"
        MOTION.apply_native_animations(
            slide,
            [{"id": "card", "kind": "fade"}, {"id": "copy", "kind": "fade"}],
            {"card": card, "copy": text},
        )
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "text-target.pptx"
            presentation.save(output)
            with ZipFile(output) as archive:
                root = etree.fromstring(archive.read("ppt/slides/slide1.xml"))
        ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
        self.assertEqual(root.xpath("count(.//p:spTgt[@spid=$card])", namespaces=ns,
                                    card=str(card.shape_id)), 2.0)
        self.assertEqual(root.xpath("count(.//p:spTgt[@spid=$text])", namespaces=ns,
                                    text=str(text.shape_id)), 2.0)
        self.assertEqual(root.xpath("count(.//p:bldP[@spid=$card and @animBg='1'])", namespaces=ns,
                                    card=str(card.shape_id)), 1.0)
        self.assertEqual(root.xpath("count(.//p:bldP[@spid=$text and @animBg])", namespaces=ns,
                                    text=str(text.shape_id)), 0.0)

    def test_motion_audit_rejects_text_background_only_builds(self):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, 1000000, 1000000)
        text = slide.shapes.add_textbox(100000, 100000, 800000, 400000)
        text.text_frame.text = "Animated text"
        MOTION.apply_native_animations(
            slide,
            [{"id": "card", "kind": "fade"}, {"id": "copy", "kind": "fade"}],
            {"card": card, "copy": text},
        )
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "audit.pptx"
            presentation.save(output)
            report = MOTION_AUDIT.audit_motion(output, include_powerpoint_com=False)
            self.assertEqual(report["result"], "pass")
            self.assertEqual(report["slides"][0]["textBackgroundOnlyBuilds"], 0)

            mutated = Path(directory) / "audit-mutated.pptx"
            with ZipFile(output) as source, ZipFile(mutated, "w") as target:
                for name in source.namelist():
                    payload = source.read(name)
                    if name == "ppt/slides/slide1.xml":
                        root = etree.fromstring(payload)
                        build = root.xpath(".//p:bldP[@spid=$text]", namespaces={
                            "p": "http://schemas.openxmlformats.org/presentationml/2006/main"
                        }, text=str(text.shape_id))[0]
                        build.set("animBg", "1")
                        payload = etree.tostring(root, xml_declaration=True, encoding="UTF-8")
                    target.writestr(name, payload)
            report = MOTION_AUDIT.audit_motion(mutated, include_powerpoint_com=False)
        self.assertEqual(report["result"], "fail")
        self.assertEqual(report["slides"][0]["textBackgroundOnlyBuilds"], 1)

    def test_measure_captures_css_motion_before_static_freeze(self):
        try:
            from playwright.sync_api import sync_playwright  # noqa: F401
        except ImportError as exc:
            self.skipTest(f"Playwright unavailable: {exc}")
        html = """<!doctype html><style>
        .slide { width: 1920px; height: 1080px; position: relative; }
        @keyframes rise { from { opacity: 0; transform: translateY(40px); } to { opacity: 1; transform: none; } }
        h1 { animation: rise 700ms 150ms both; }
        </style><section class='slide' data-pptx-slide><h1>Animated title</h1></section>"""
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "motion.html"
            output = Path(directory) / "motion.pptx"
            source.write_text(html, encoding="utf-8")
            data = MEASURE.measure(source, no_screenshots=True, verbose=False)
            ASSEMBLE.assemble(data, output)
            with ZipFile(output) as archive:
                root = etree.fromstring(archive.read("ppt/slides/slide1.xml"))
        slide = data["slides"][0]
        self.assertEqual(slide["motions"][0]["kind"], "fade")
        self.assertEqual(slide["motions"][0]["delay"], 150)
        self.assertTrue(any(record.get("motionId") == slide["motions"][0]["id"]
                            for record in slide["records"]))
        ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
        self.assertEqual(root.xpath("count(.//p:animEffect[@filter='fade'])", namespaces=ns), 1.0)

    def test_measure_captures_explicit_motion_plan(self):
        try:
            from playwright.sync_api import sync_playwright  # noqa: F401
        except ImportError as exc:
            self.skipTest(f"Playwright unavailable: {exc}")
        html = """<!doctype html><style>
        .slide { width: 1920px; height: 1080px; position: relative; }
        </style><section class='slide' data-pptx-slide>
        <div data-pptx-motion-plan='[{"effect":"fade","duration":300,"trigger":"click"},{"effect":"motion","path":"M 0 0 L 0.2 0 E","duration":500,"trigger":"with"}]'>Complex</div>
        </section>"""
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "plan.html"
            source.write_text(html, encoding="utf-8")
            data = MEASURE.measure(source, no_screenshots=True, verbose=False)
        plan = data["slides"][0]["motions"][0]["plan"]
        self.assertEqual([item["effect"] for item in plan], ["fade", "motion"])
        self.assertEqual(plan[1]["trigger"], "with")

    def test_container_motion_groups_descendant_shapes_on_one_click(self):
        try:
            from playwright.sync_api import sync_playwright  # noqa: F401
        except ImportError as exc:
            self.skipTest(f"Playwright unavailable: {exc}")
        html = """<!doctype html><style>
        .slide { width: 1920px; height: 1080px; position: relative; }
        .group { animation: reveal 500ms both; }
        .card { display: inline-block; width: 360px; height: 180px; margin: 20px;
                background: #d8eee9; border-top: 12px solid #123247; }
        @keyframes reveal { from { opacity: 0; } to { opacity: 1; } }
        </style><section class='slide' data-pptx-slide><div class='group'
        data-pptx-motion-plan='[{"effect":"fade","duration":500,"trigger":"click"}]'>
        <div class='card'>First card</div><div class='card'>Second card</div></div></section>"""
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "group.html"
            output = Path(directory) / "group.pptx"
            source.write_text(html, encoding="utf-8")
            data = MEASURE.measure(source, no_screenshots=True, verbose=False)
            ASSEMBLE.assemble(data, output)
            with ZipFile(output) as archive:
                root = etree.fromstring(archive.read("ppt/slides/slide1.xml"))
        slide = data["slides"][0]
        motion_id = slide["motions"][0]["id"]
        self.assertGreaterEqual(sum(record.get("motionId") == motion_id for record in slide["records"]), 2)
        ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
        self.assertGreaterEqual(root.xpath("count(.//p:animEffect[@filter='fade'])", namespaces=ns), 2.0)
        self.assertGreaterEqual(root.xpath("count(.//p:cTn[@nodeType='withEffect'])", namespaces=ns), 1.0)

    def test_measure_decomposes_parallel_css_transforms(self):
        try:
            from playwright.sync_api import sync_playwright  # noqa: F401
        except ImportError as exc:
            self.skipTest(f"Playwright unavailable: {exc}")
        html = """<!doctype html><style>
        .slide { width: 1920px; height: 1080px; position: relative; }
        @keyframes reveal {
          from { opacity: 0; transform: translate(40px, 20px) scale(.8) rotate(-20deg); }
          to { opacity: 1; transform: none; }
        }
        .hero { position: absolute; left: 100px; top: 100px; animation: reveal 600ms both; }
        </style><section class='slide' data-pptx-slide><div class='hero'>Hero</div></section>"""
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "parallel.html"
            source.write_text(html, encoding="utf-8")
            data = MEASURE.measure(source, no_screenshots=True, verbose=False)
        plan = data["slides"][0]["motions"][0]["plan"]
        self.assertEqual([item["effect"] for item in plan], ["fade", "motion", "scale", "rotate"])
        self.assertEqual([item["trigger"] for item in plan], ["click", "with", "with", "with"])
        self.assertEqual(plan[1]["path"], "M 0.020833 0.018519 L 0 0 E")
        self.assertAlmostEqual(plan[2]["from"][0], 80.0, places=4)
        self.assertAlmostEqual(plan[2]["from"][1], 80.0, places=4)
        self.assertEqual(plan[3]["from"], -20)

    @unittest.skipUnless(os.name == "nt", "PowerPoint COM validation is Windows-only")
    def test_powerpoint_recognizes_emitted_fade(self):
        try:
            import win32com.client
        except ImportError as exc:
            self.skipTest(f"pywin32 unavailable: {exc}")
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "powerpoint-motion.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            fade_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, 1000000, 1000000)
            spin_shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, 1000000, 0, 1000000, 1000000)
            MOTION.apply_native_animations(
                slide,
                [
                    {"id": "fade", "kind": "fade", "duration": 600, "delay": 0},
                    {"id": "spin", "kind": "spin", "duration": 800, "delay": 0},
                ],
                {"fade": fade_shape, "spin": spin_shape},
            )
            presentation.save(output)
            app = win32com.client.Dispatch("PowerPoint.Application")
            try:
                document = app.Presentations.Open(str(output), False, True, False)
                try:
                    sequence = document.Slides.Item(1).TimeLine.MainSequence
                    self.assertEqual(sequence.Count, 2)
                    self.assertEqual(sequence.Item(1).EffectType, 10)
                    self.assertEqual(sequence.Item(2).EffectType, 61)
                finally:
                    document.Close()
            finally:
                app.Quit()

    @unittest.skipUnless(os.name == "nt", "PowerPoint COM validation is Windows-only")
    def test_powerpoint_recognizes_text_target(self):
        try:
            import win32com.client
        except ImportError as exc:
            self.skipTest(f"pywin32 unavailable: {exc}")
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "powerpoint-text-target.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            text = slide.shapes.add_textbox(0, 0, 1000000, 400000)
            text.text_frame.text = "Animated text"
            MOTION.apply_native_animations(
                slide, [{"id": "copy", "kind": "fade", "duration": 600}], {"copy": text}
            )
            presentation.save(output)
            app = win32com.client.Dispatch("PowerPoint.Application")
            try:
                document = app.Presentations.Open(str(output), False, True, False)
                try:
                    sequence = document.Slides.Item(1).TimeLine.MainSequence
                    self.assertEqual(sequence.Count, 1)
                    self.assertEqual(sequence.Item(1).EffectType, 10)
                    self.assertEqual(sequence.Item(1).Shape.TextFrame.TextRange.Text, "Animated text")
                finally:
                    document.Close()
            finally:
                app.Quit()

    @unittest.skipUnless(os.name == "nt", "PowerPoint COM validation is Windows-only")
    def test_powerpoint_recognizes_composable_motion_plan(self):
        try:
            import win32com.client
        except ImportError as exc:
            self.skipTest(f"pywin32 unavailable: {exc}")
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "powerpoint-complex-motion.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, 1000000, 1000000)
            MOTION.apply_native_animations(slide, [{
                "id": "hero",
                "plan": [
                    {"effect": "fade", "duration": 300, "trigger": "click"},
                    {"effect": "motion", "path": "M 0 0 L 0.18 -0.08 E", "duration": 650, "trigger": "with"},
                    {"effect": "scale", "from": [80, 80], "to": [115, 115], "duration": 650, "trigger": "with"},
                    {"effect": "rotate", "by": 360, "duration": 650, "trigger": "with"},
                    {"effect": "fade", "duration": 250, "trigger": "after"},
                ],
            }], {"hero": shape})
            presentation.save(output)
            app = win32com.client.Dispatch("PowerPoint.Application")
            try:
                document = app.Presentations.Open(str(output), False, True, False)
                try:
                    sequence = document.Slides.Item(1).TimeLine.MainSequence
                    self.assertEqual(sequence.Count, 5)
                    self.assertEqual([sequence.Item(i).Timing.TriggerType for i in range(1, 6)], [1, 2, 2, 2, 3])
                    self.assertEqual(sequence.Item(2).EffectType, 0)
                    self.assertEqual(sequence.Item(3).EffectType, 0)
                    self.assertEqual(sequence.Item(4).EffectType, 61)
                finally:
                    document.Close()
            finally:
                app.Quit()


if __name__ == "__main__":
    unittest.main()
